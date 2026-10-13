"""Load the workbook and create an auditable canonical data table."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pandas as pd


TARGET_ORIGINAL = "Y 最低中性Li空位形成能 E_f(V_Li^0) (eV)"

FEATURE_SPECS: dict[str, dict[str, str]] = {
    "核心1 Formation energy / atom, Ef (eV/atom)": {
        "canonical": "formation_energy_per_atom",
        "group": "T",
        "subgroup": "thermodynamic",
        "unit": "eV/atom",
        "meaning": "Bulk formation energy per atom",
        "source": "DFT / Materials Project",
    },
    "核心2 Li平均Bader电荷, qLi (e)": {
        "canonical": "li_mean_bader_charge",
        "group": "E",
        "subgroup": "Bader",
        "unit": "e",
        "meaning": "Mean Li Bader charge loss qLi",
        "source": "DFT charge density + Bader",
    },
    "核心3 Li平均配位数, CN_Li": {
        "canonical": "li_mean_coordination_number",
        "group": "G",
        "subgroup": "coordination",
        "unit": "dimensionless",
        "meaning": "Mean Li coordination number",
        "source": "CIF structure analysis",
    },
    "核心4 归一化Li–X距离, d_norm": {
        "canonical": "normalized_li_x_distance",
        "group": "G",
        "subgroup": "bond geometry",
        "unit": "dimensionless",
        "meaning": "Li-X distance divided by the sum of tabulated radii",
        "source": "CIF + atomic/ionic radius table",
    },
    "核心5 每Li对应体积, V/N_Li (Å³/Li)": {
        "canonical": "volume_per_li",
        "group": "G",
        "subgroup": "volume",
        "unit": "Å^3/Li",
        "meaning": "Unit-cell volume per Li atom",
        "source": "CIF-derived",
    },
    "核心6 Li–骨架电负性差, Δχ": {
        "canonical": "li_framework_electronegativity_difference",
        "group": "C",
        "subgroup": "electronegativity",
        "unit": "Pauling scale",
        "meaning": "Framework mean electronegativity minus Li electronegativity",
        "source": "Formula/CIF + electronegativity table",
    },
    "核心7 Band gap, Eg (eV)": {
        "canonical": "band_gap",
        "group": "E",
        "subgroup": "band gap",
        "unit": "eV",
        "meaning": "Electronic band gap",
        "source": "DFT / Materials Project",
    },
    "候选1 Energy above hull, E_hull (eV/atom)": {
        "canonical": "energy_above_hull",
        "group": "T",
        "subgroup": "phase stability",
        "unit": "eV/atom",
        "meaning": "Energy above the competing-phase hull",
        "source": "DFT database + phase-diagram postprocessing / MP",
    },
    "候选2 Li atomic fraction, x_Li": {
        "canonical": "li_atomic_fraction",
        "group": "C",
        "subgroup": "stoichiometry",
        "unit": "fraction",
        "meaning": "Li atoms divided by total atoms",
        "source": "Chemical formula / CIF",
    },
    "候选3 Density, ρ (g/cm³)": {
        "canonical": "density",
        "group": "G",
        "subgroup": "bulk geometry",
        "unit": "g/cm^3",
        "meaning": "Crystal density",
        "source": "CIF-derived / MP",
    },
    "候选4 Volume / atom, V/N_atom (Å³/atom)": {
        "canonical": "volume_per_atom",
        "group": "G",
        "subgroup": "volume",
        "unit": "Å^3/atom",
        "meaning": "Unit-cell volume per atom",
        "source": "CIF-derived",
    },
    "候选5 平均Li–X键长, d̄_Li-X (Å)": {
        "canonical": "mean_li_x_bond_length",
        "group": "G",
        "subgroup": "bond geometry",
        "unit": "Å",
        "meaning": "Mean nearest-neighbor Li-X bond length",
        "source": "CIF structure analysis",
    },
    "候选6 Li–X键长标准差, σ_d (Å)": {
        "canonical": "li_x_bond_length_std",
        "group": "G",
        "subgroup": "bond geometry",
        "unit": "Å",
        "meaning": "Standard deviation of nearest-neighbor Li-X distances",
        "source": "CIF structure analysis",
    },
    "候选7 Li–X键长范围, Δd (Å)": {
        "canonical": "li_x_bond_length_range",
        "group": "G",
        "subgroup": "bond geometry",
        "unit": "Å",
        "meaning": "Maximum minus minimum nearest-neighbor Li-X distance",
        "source": "CIF structure analysis",
    },
    "候选8 Li配位数标准差, σ_CN": {
        "canonical": "li_coordination_number_std",
        "group": "G",
        "subgroup": "coordination",
        "unit": "dimensionless",
        "meaning": "Standard deviation of Li coordination numbers",
        "source": "CIF structure analysis",
    },
    "候选9 平均框架电负性, χ̄_framework": {
        "canonical": "mean_framework_electronegativity",
        "group": "C",
        "subgroup": "electronegativity",
        "unit": "Pauling scale",
        "meaning": "Framework electronegativity weighted by non-Li atom fraction",
        "source": "Chemical formula + electronegativity table",
    },
    "候选10 平均框架原子半径, r̄_framework": {
        "canonical": "mean_framework_atomic_radius",
        "group": "C",
        "subgroup": "composition",
        "unit": "Å",
        "meaning": "Mean framework atomic radius",
        "source": "Chemical formula + elemental radius table",
    },
    "候选11 平均原子序数, Z̄": {
        "canonical": "mean_atomic_number",
        "group": "C",
        "subgroup": "composition",
        "unit": "dimensionless",
        "meaning": "Mean atomic number over all atoms",
        "source": "Chemical formula + periodic table",
    },
    "候选12 平均原子质量, M̄ (g/mol-atom)": {
        "canonical": "mean_atomic_mass",
        "group": "C",
        "subgroup": "composition",
        "unit": "g/mol-atom",
        "meaning": "Mean atomic mass over all atoms",
        "source": "Chemical formula + atomic-mass table",
    },
    "候选13 Bulk平均ELF (3D)": {
        "canonical": "bulk_mean_elf",
        "group": "E",
        "subgroup": "ELF",
        "unit": "dimensionless",
        "meaning": "Mean three-dimensional bulk ELF",
        "source": "Bulk ELFCAR grid",
    },
    "候选14 ELF≥0.8体积分数 (%)": {
        "canonical": "elf_high_fraction",
        "group": "E",
        "subgroup": "ELF",
        "unit": "%",
        "meaning": "Volume fraction with ELF >= 0.8",
        "source": "Bulk ELFCAR grid",
    },
    "候选15 ELF≤0.2体积分数 (%)": {
        "canonical": "elf_low_fraction",
        "group": "E",
        "subgroup": "ELF",
        "unit": "%",
        "meaning": "Volume fraction with ELF <= 0.2",
        "source": "Bulk ELFCAR grid",
    },
    "候选16 骨架电负性标准差 σ_χ": {
        "canonical": "framework_electronegativity_std",
        "group": "C",
        "subgroup": "electronegativity",
        "unit": "Pauling scale",
        "meaning": "Population standard deviation of framework electronegativities",
        "source": "Reconstructed from formula and Pauling electronegativities",
    },
    "候选17 框架含氧原子分数 xO_fw": {
        "canonical": "framework_oxygen_fraction",
        "group": "C",
        "subgroup": "stoichiometry",
        "unit": "fraction",
        "meaning": "O atoms divided by all non-Li framework atoms",
        "source": "Reconstructed from chemical formula",
    },
    "候选18 空位浓度": {
        "canonical": "vacancy_concentration",
        "group": "AUX",
        "subgroup": "defect-model control",
        "unit": "%",
        "meaning": "Nominal Li-vacancy concentration used in defect calculations",
        "source": "Defect supercell setup",
    },
    "候选19 Li_mean_charge_loss": {
        "canonical": "li_mean_charge_loss",
        "group": "E",
        "subgroup": "Bader duplicate",
        "unit": "e",
        "meaning": "Duplicate representation of mean Li Bader charge loss",
        "source": "Workbook-derived duplicate",
    },
    "候选20 vacancy_concentration_percent": {
        "canonical": "vacancy_concentration_percent",
        "group": "AUX",
        "subgroup": "defect-model control duplicate",
        "unit": "%",
        "meaning": "Duplicate representation of nominal vacancy concentration",
        "source": "Workbook-derived duplicate",
    },
}

TARGET_CANONICAL = "Li_vacancy_formation_energy"

PAULING_ELECTRONEGATIVITY: dict[str, float] = {
    "Li": 0.98,
    "B": 2.04,
    "C": 2.55,
    "N": 3.04,
    "O": 3.44,
    "F": 3.98,
    "Si": 1.90,
    "Al": 1.61,
    "P": 2.19,
    "S": 2.58,
    "Se": 2.55,
    "Fe": 1.83,
    "Co": 1.88,
    "Ni": 1.91,
    "Cu": 1.90,
    "Mn": 1.55,
    "Zn": 1.65,
    "Mo": 2.16,
    "Zr": 1.33,
}

ELECTRONIC_FEATURES = [
    "band_gap",
    "li_mean_bader_charge",
    "bulk_mean_elf",
    "elf_high_fraction",
    "elf_low_fraction",
]
GEOMETRIC_FEATURES = [
    "li_mean_coordination_number",
    "normalized_li_x_distance",
    "volume_per_li",
    "density",
    "volume_per_atom",
    "mean_li_x_bond_length",
    "li_x_bond_length_std",
    "li_x_bond_length_range",
    "li_coordination_number_std",
]
CHEMICAL_FEATURES = [
    "li_atomic_fraction",
    "li_framework_electronegativity_difference",
    "mean_framework_atomic_radius",
    "mean_atomic_number",
    "mean_atomic_mass",
    "framework_electronegativity_std",
    "framework_oxygen_fraction",
]
THERMODYNAMIC_FEATURES = ["formation_energy_per_atom", "energy_above_hull"]
AUXILIARY_FEATURES = ["vacancy_concentration"]

# Exact duplicates and an exact affine re-expression are kept in the audit but are
# excluded from the non-redundant main feature matrix.
REDUNDANT_FEATURES = [
    "mean_framework_electronegativity",
    "li_mean_charge_loss",
    "vacancy_concentration_percent",
]
ALL_INTRINSIC_FEATURES = (
    ELECTRONIC_FEATURES
    + GEOMETRIC_FEATURES
    + CHEMICAL_FEATURES
    + THERMODYNAMIC_FEATURES
)
ALL_VALID_FEATURES = ALL_INTRINSIC_FEATURES + AUXILIARY_FEATURES


def parse_formula(formula: str) -> dict[str, int]:
    """Parse the simple reduced chemical formulas used in the workbook."""

    cleaned = str(formula).strip().replace("(", "").replace(")", "")
    tokens = re.findall(r"([A-Z][a-z]?)(\d*)", cleaned)
    if not tokens or "".join(element + count for element, count in tokens) != cleaned:
        raise ValueError(f"Unsupported chemical formula: {formula!r}")
    counts: dict[str, int] = {}
    for element, count in tokens:
        counts[element] = counts.get(element, 0) + int(count or 1)
    return counts


def reconstruct_formula_descriptors(formula: str) -> dict[str, Any]:
    counts = parse_formula(formula)
    framework = {element: count for element, count in counts.items() if element != "Li"}
    framework_total = sum(framework.values())
    unknown = sorted(set(framework) - set(PAULING_ELECTRONEGATIVITY))
    if not framework_total or unknown:
        raise ValueError(f"Cannot reconstruct framework descriptors for {formula}: {unknown}")
    values = np.array(
        [PAULING_ELECTRONEGATIVITY[element] for element, count in framework.items() for _ in range(count)],
        dtype=float,
    )
    return {
        "formula_counts": counts,
        "framework_total_atoms": framework_total,
        "reconstructed_mean_framework_electronegativity": float(values.mean()),
        "reconstructed_framework_electronegativity_std": float(values.std(ddof=0)),
        "reconstructed_framework_oxygen_fraction": float(framework.get("O", 0) / framework_total),
    }


def _read_formula_metadata(excel_path: Path) -> tuple[pd.DataFrame, dict[str, Any]]:
    workbook_values = pd.read_excel(
        excel_path,
        sheet_name="24种材料特征矩阵",
        engine="openpyxl",
    )
    workbook_formula = openpyxl.load_workbook(excel_path, data_only=False, read_only=True)
    workbook_cached = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
    ws_formula = workbook_formula["24种材料特征矩阵"]
    ws_cached = workbook_cached["24种材料特征矩阵"]
    headers = [cell.value for cell in ws_formula[1]]
    formulas: dict[str, list[Any]] = {str(header): [] for header in headers}
    cached_values: dict[str, list[Any]] = {str(header): [] for header in headers}
    for row in range(2, ws_formula.max_row + 1):
        for col, header in enumerate(headers, start=1):
            formulas[str(header)].append(ws_formula.cell(row, col).value)
            cached_values[str(header)].append(ws_cached.cell(row, col).value)
    external_links = []
    for link in workbook_formula._external_links:
        target = getattr(getattr(link, "file_link", None), "Target", None)
        if target is None:
            target = getattr(getattr(link, "externalBook", None), "id", None)
        external_links.append(str(target or "external workbook"))
    metadata = {
        "formulas": formulas,
        "cached_values": cached_values,
        "external_links": external_links,
        "sheet_names": workbook_formula.sheetnames,
        "n_rows": ws_formula.max_row - 1,
        "n_columns": ws_formula.max_column,
    }
    workbook_formula.close()
    workbook_cached.close()
    return workbook_values, metadata


def load_dataset(excel_path: str | Path) -> dict[str, Any]:
    """Load and repair the workbook without changing the input file."""

    excel_path = Path(excel_path)
    raw, workbook_metadata = _read_formula_metadata(excel_path)
    required = ["Material", "来源", TARGET_ORIGINAL, *FEATURE_SPECS.keys()]
    missing = [column for column in required if column not in raw.columns]
    if missing:
        raise KeyError(f"Missing required workbook columns: {missing}")

    canonical = pd.DataFrame(index=raw.index)
    canonical["material"] = raw["Material"].astype(str).str.strip()
    canonical["source"] = raw["来源"].astype(str).str.strip()
    canonical["mp_id"] = canonical["source"].str.extract(r"(mp-\d+)", expand=False)
    canonical["chemical_formula"] = canonical["material"]
    canonical[TARGET_CANONICAL] = pd.to_numeric(raw[TARGET_ORIGINAL], errors="coerce")

    reconstruction_rows: list[dict[str, Any]] = []
    for formula in canonical["chemical_formula"]:
        reconstruction_rows.append(reconstruct_formula_descriptors(formula))

    for original, spec in FEATURE_SPECS.items():
        canonical_name = spec["canonical"]
        canonical[canonical_name] = pd.to_numeric(raw[original], errors="coerce")

    reconstruction = pd.DataFrame(reconstruction_rows, index=canonical.index)
    canonical["framework_electronegativity_std"] = reconstruction[
        "reconstructed_framework_electronegativity_std"
    ]
    canonical["framework_oxygen_fraction"] = reconstruction[
        "reconstructed_framework_oxygen_fraction"
    ]
    canonical["reconstructed_mean_framework_electronegativity"] = reconstruction[
        "reconstructed_mean_framework_electronegativity"
    ]
    canonical["framework_total_atoms"] = reconstruction["framework_total_atoms"]

    raw_sigma = pd.to_numeric(raw["候选16 骨架电负性标准差 σ_χ"], errors="coerce")
    raw_oxygen = pd.to_numeric(raw["候选17 框架含氧原子分数 xO_fw"], errors="coerce")
    canonical["raw_formula_sigma_chi"] = raw_sigma
    canonical["raw_formula_xO_fw"] = raw_oxygen
    canonical["sigma_chi_reconstruction_abs_diff"] = (
        raw_sigma - canonical["framework_electronegativity_std"]
    ).abs()
    canonical["xO_reconstruction_abs_diff"] = (
        raw_oxygen - canonical["framework_oxygen_fraction"]
    ).abs()
    canonical["category"] = canonical["chemical_formula"].map(assign_chemical_category)

    formula_flags = []
    for index, formula in enumerate(canonical["chemical_formula"]):
        formula_value = workbook_metadata["formulas"]["候选16 骨架电负性标准差 σ_χ"][index]
        oxygen_value = workbook_metadata["formulas"]["候选17 框架含氧原子分数 xO_fw"][index]
        formula_flags.append(
            {
                "sigma_formula": formula_value,
                "oxygen_formula": oxygen_value,
                "sigma_formula_external": isinstance(formula_value, str) and "[1]" in formula_value,
                "oxygen_formula_external": isinstance(oxygen_value, str) and "[1]" in oxygen_value,
                "sigma_formula_is_name_error": str(formula_value).strip() == "#NAME?",
                "oxygen_formula_is_name_error": str(oxygen_value).strip() == "#NAME?",
            }
        )
    formula_flags_df = pd.DataFrame(formula_flags, index=canonical.index)
    canonical = pd.concat([canonical, formula_flags_df], axis=1)

    return {
        "raw": raw,
        "canonical": canonical,
        "workbook_metadata": workbook_metadata,
        "feature_specs": FEATURE_SPECS,
        "target_original": TARGET_ORIGINAL,
    }


def assign_chemical_category(formula: str) -> str:
    """Assign one transparent, mutually exclusive chemical category."""

    elements = set(parse_formula(formula))
    if "C" in elements:
        return "carbon-containing"
    if elements & {"F", "Cl", "Br", "I"}:
        return "halide"
    if elements & {"S", "Se", "Te"}:
        return "sulfide-selenide"
    if "P" in elements:
        return "phosphide/phosphate"
    if "N" in elements:
        return "nitride/nitrate"
    if "O" in elements:
        return "oxygen-containing"
    return "other"


def feature_dictionary_frame() -> pd.DataFrame:
    rows = []
    for original, spec in FEATURE_SPECS.items():
        canonical = spec["canonical"]
        status = "redundant-audit-only" if canonical in REDUNDANT_FEATURES else "main-available"
        note = ""
        if canonical == "framework_electronegativity_std":
            note = "Reconstructed independently; cached workbook values agree within floating-point tolerance."
        elif canonical == "framework_oxygen_fraction":
            note = "Reconstructed independently because workbook external formula duplicates sigma_chi."
        elif canonical == "li_mean_charge_loss":
            note = "Compared with li_mean_bader_charge; retained for audit, excluded from main matrix."
        elif canonical == "vacancy_concentration_percent":
            note = "Compared with vacancy_concentration; retained for audit, excluded from main matrix."
        elif canonical == "mean_framework_electronegativity":
            note = "Affine transform of li_framework_electronegativity_difference using χ(Li)=0.98; excluded from main matrix."
        rows.append(
            {
                "canonical_name": canonical,
                "original_column": original,
                "group": spec["group"],
                "subgroup": spec["subgroup"],
                "unit": spec["unit"],
                "physical_meaning": spec["meaning"],
                "source": spec["source"],
                "status": status,
                "notes": note,
            }
        )
    rows.append(
        {
            "canonical_name": TARGET_CANONICAL,
            "original_column": TARGET_ORIGINAL,
            "group": "Y",
            "subgroup": "target",
            "unit": "eV",
            "physical_meaning": "Minimum neutral Li-vacancy formation energy over all inequivalent Li sites",
            "source": "DFT defect calculations; 46 sites reduced to 24 material-level minima",
            "status": "target",
            "notes": "Not used as a feature or transformed target.",
        }
    )
    return pd.DataFrame(rows)
